import pytest
import pypdf
import docx
import yaml
from memoria.core.chunker import Chunker, parse_document, SUPPORTED


def test_supported_extensions():
    for ext in [".md", ".markdown", ".txt", ".pdf", ".docx", ".yaml", ".yml"]:
        assert ext in SUPPORTED


def test_parse_document_markdown(tmp_path):
    f = tmp_path / "test.md"
    f.write_text("# Title\nContent here.", encoding="utf-8")
    assert "Title" in parse_document(str(f))
    chunks = Chunker().split(str(f))
    assert len(chunks) >= 1
    assert "Title" in chunks[0]


def test_parse_document_pdf(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    writer = pypdf.PdfWriter()
    p = writer.add_blank_page(width=200, height=200)
    # Write a minimal page with text stream
    from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 150 Td (Hello PDF World) Tj ET")
    p[NameObject("/Contents")] = stream
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    p[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): fonts})
    with open(pdf_path, "wb") as f:
        writer.write(f)

    parsed = parse_document(str(pdf_path))
    assert "Hello PDF World" in parsed
    chunks = Chunker().split(str(pdf_path))
    assert len(chunks) >= 1
    assert "Hello PDF World" in chunks[0]


def test_parse_document_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Docx Heading", 0)
    doc.add_paragraph("First paragraph text.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header1"
    table.cell(0, 1).text = "Header2"
    table.cell(1, 0).text = "Value1"
    table.cell(1, 1).text = "Value2"
    doc.add_paragraph("Second paragraph text.")
    doc.save(str(docx_path))

    parsed = parse_document(str(docx_path))
    assert "Docx Heading" in parsed
    assert "First paragraph text." in parsed
    assert "Header1 | Header2" in parsed
    assert "Value1 | Value2" in parsed
    assert "Second paragraph text." in parsed

    chunks = Chunker().split(str(docx_path))
    assert len(chunks) >= 1
    joined = " ".join(chunks)
    assert "Docx Heading" in joined
    assert "Header1 | Header2" in joined


def test_parse_document_yaml(tmp_path):
    yaml_path = tmp_path / "config.yaml"
    yaml_path.write_text("server:\n  host: 127.0.0.1\n  port: 8000\n", encoding="utf-8")

    parsed = parse_document(str(yaml_path))
    assert "server:" in parsed
    assert "127.0.0.1" in parsed

    chunks = Chunker().split(str(yaml_path))
    assert len(chunks) >= 1
    assert "127.0.0.1" in chunks[0]


def test_parse_empty_document(tmp_path):
    f = tmp_path / "empty.txt"
    f.write_text("", encoding="utf-8")
    assert parse_document(str(f)) == ""
    assert Chunker().split(str(f)) == []
